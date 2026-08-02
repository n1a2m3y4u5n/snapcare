"""
proposal_ko.md → 대회 제출용 .docx.

공식 템플릿을 열어 그 안의 빈 칸을 채운다. 새 문서를 만들지 않는 이유는 표지·안내문·
쪽 설정이 심사 서식 그 자체이기 때문이다 — 우리가 다시 그리면 서식이 달라진다.

본문은 마크다운의 아주 좁은 부분집합만 쓴다(제목, 문단, 표, 그림, 코드블록,
굵게/기울임/코드). 범용 변환기를 쓰지 않는 것은 pandoc 도 LibreOffice 도 없는
환경에서 돌아야 하고, 표 22개와 그림 14장이 서식 안에서 어떻게 앉는지를
직접 정해야 하기 때문이다.

    python3 docs/proposal/tools/build_docx.py [--out 경로]
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
DOC = HERE.parent                      # docs/proposal
LANG = 'KO'                            # main() 에서 원본 파일 이름으로 정한다
ROOT = HERE.parents[2]                 # 레포 루트

# 공식 서식은 레포에 함께 둔다. 이것이 없으면 다른 사람이 문서를 다시 만들 수 없고,
# 심사 서식이 남의 다운로드 폴더에만 있는 상태가 된다.
TEMPLATE = DOC / 'template.docx'

# 본문이 앉을 칸. (표 번호, 행 번호) — 템플릿의 "Content" 라벨 바로 아래 빈 행이다.
SLOTS = {
    'Summary of Web App':            (0, 9),
    'Problem Statement and Background': (0, 12),
    'Solution and Main Functions':   (0, 15),
    'User Flow of the Web App':      (1, 2),
    'Development and Implementation': (1, 5),
    'Results and Expected Impact':   (1, 8),
}

HEADER = {                             # 표 0 위쪽 정보칸
    (1, 1): '[H1] SnapCare',
    (1, 3): 'SnapCare',
    (2, 1): 'Yunsu NAM',
    (2, 3): 'Cristian RUDICO',
    (3, 1): 'Dohun KIM',
    (3, 3): 'Daniel STANCHEVICI',
    (4, 1): ('Target 3.4 Noncommunicable diseases · Target 3.8 Universal health coverage · '
             'Target 3.d Health risk management'),
    (5, 1): 'Github Repository: https://github.com/16th-eicon-contest/h1-SnapCare',
    (6, 1): 'Deployment URL: https://snapcare.fly.dev',
}

# 영어판만 덮어쓰는 칸. 로마자를 앞에 둔다. 한글을 지우지는 않는다. 본인 이름이다.
# 이름은 이제 두 판이 같으므로 덮어쓸 것이 없다. 빈 사전을 남겨 두는 것은
# 앞으로 영어판에서만 달라질 칸이 생길 자리이기 때문이다.
HEADER_EN: dict = {}

BODY_PT = 10        # 세리프는 같은 pt 에서 작아 보인다
LINE = 1.15         # 논문 본문에 가까운 행간
MAX_W, MAX_H = 6.0, 5.2                # 그림 최대 크기(인치)
MONO = 'Consolas'
# 논문 조판의 관례를 따른다 — 라틴은 Times New Roman, 한글은 바탕(명조 계열).
# 바탕을 고른 것은 심사자가 Windows 워드로 열 가능성이 높기 때문이다. macOS 에
# 바탕이 없으면 워드가 다른 명조로 대체하는데, 그래도 세리프라는 성격은 남는다.
LATIN, HANGUL = 'Times New Roman', '바탕'
MUTED = RGBColor(0x60, 0x60, 0x68)


# ── 문서 조작 도우미 ────────────────────────────────────────────────

def clear(cell):
    """칸을 비운다. 템플릿의 빈 문단 하나가 남아 있으면 앞머리에 빈 줄이 생긴다."""
    for p in list(cell.paragraphs)[1:]:
        p._element.getparent().remove(p._element)
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    return p


def borders(table):
    """표 테두리. 템플릿에 'Table Grid' 스타일이 없어 직접 그린다."""
    el = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{side}')
        e.set(qn('w:val'), 'single')
        e.set(qn('w:sz'), '4')
        e.set(qn('w:color'), 'BFBFBF')
        el.append(e)
    table._tbl.tblPr.append(el)


def set_font(run, name=None):
    """라틴과 한글 글꼴을 함께 지정한다.

    python-docx 의 font.name 은 라틴 글꼴만 바꾼다. 한글은 rFonts 의 eastAsia
    속성이 따로 결정하고, 비워 두면 Word 가 알아서 명조 계열로 떨어뜨려
    템플릿의 안내문과 본문이 다른 글꼴로 보인다.
    """
    run.font.name = name or LATIN
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), MONO if name == MONO else HANGUL)


def shade(cell, hex_color):
    e = OxmlElement('w:shd')
    e.set(qn('w:val'), 'clear')
    e.set(qn('w:fill'), hex_color)
    cell._tc.get_or_add_tcPr().append(e)


INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`|\*[^*]+?\*)')


def write_inline(par, text, size=BODY_PT, bold=False):
    """**굵게**, `코드`, *기울임* 만 처리한다. 본문이 쓰는 것이 이것뿐이다."""
    text = text.replace('\\<', '<').replace('\\>', '>')
    for piece in INLINE.split(text):
        if not piece:
            continue
        r = par.add_run()
        if piece.startswith('**') and piece.endswith('**'):
            r.text, r.bold = piece[2:-2], True
            set_font(r)
        elif piece.startswith('`') and piece.endswith('`'):
            r.text = piece[1:-1]
            set_font(r, MONO)
            r.font.size = Pt(size - 0.5)
        elif piece.startswith('*') and piece.endswith('*'):
            r.text, r.italic = piece[1:-1], True
            set_font(r)
        else:
            r.text = piece
            set_font(r)
        r.font.size = r.font.size or Pt(size)
        if bold:
            r.bold = True


# ── 마크다운 → 블록 ────────────────────────────────────────────────

def blocks(md: str):
    """본문을 (종류, 값) 목록으로 자른다."""
    lines = md.split('\n')
    i, out = 0, []
    while i < len(lines):
        ln = lines[i]
        if not ln.strip():
            i += 1
        elif ln.startswith('<!--'):
            # 부록의 시작·끝 표지 같은 주석. 원본에서는 필요하지만 서식에
            # 그대로 찍히면 독자에게는 뜻 없는 줄이다. 실제로 네 줄이 찍혔다.
            i += 1
        elif ln.startswith('> '):
            j = i
            while j < len(lines) and lines[j].startswith('> '):
                j += 1
            out.append(('quote', [x[2:].strip() for x in lines[i:j]]))
            i = j
        elif ln.startswith('#### '):
            out.append(('h4', ln[5:].strip()))
            i += 1
        elif ln.startswith('### '):
            out.append(('h', ln[4:].strip()))
            i += 1
        elif ln.startswith('```'):
            j = i + 1
            while j < len(lines) and not lines[j].startswith('```'):
                j += 1
            out.append(('code', '\n'.join(lines[i + 1:j])))
            i = j + 1
        elif ln.startswith('!['):
            m = re.match(r'!\[[^\]]*\]\(([^)]+)\)', ln)
            out.append(('img', m.group(1)))
            i += 1
        elif ln.startswith('|'):
            j = i
            while j < len(lines) and lines[j].startswith('|'):
                j += 1
            rows = [[c.strip() for c in r.strip().strip('|').split('|')]
                    for r in lines[i:j]]
            rows = [r for r in rows if not all(set(c) <= set('-: ') for c in r)]
            out.append(('table', rows))
            i = j
        else:
            out.append(('p', ln.strip()))
            i += 1
    return out


def render(cell, md: str):
    first = clear(cell)
    used_first = False

    def para():
        nonlocal used_first
        if not used_first:
            used_first = True
            return first
        return cell.add_paragraph()

    for kind, val in blocks(md):
        if kind == 'h':
            p = para()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after = Pt(2)
            write_inline(p, val, size=BODY_PT + 1.5, bold=True)
        elif kind == 'h4':
            # 절 안의 소제목. 위 단계보다 작게 잡아야 층이 구분된다.
            p = para()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(1)
            write_inline(p, val, size=BODY_PT + 0.5, bold=True)

        elif kind == 'p':
            p = para()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = LINE
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            # 그림 설명은 가운데로, 조금 작게.
            if val.startswith('**[그림') or val.startswith('**[Figure'):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(10)
                write_inline(p, val, size=BODY_PT - 1)
                for r in p.runs:
                    r.font.color.rgb = MUTED
            else:
                write_inline(p, val)

        elif kind == 'quote':
            # 인용은 들여쓰고 회색으로. 우리 문장과 남의 문장이 구별되어야 한다.
            for k, t in enumerate(val):
                q = para()
                q.paragraph_format.left_indent = Inches(0.28)
                q.paragraph_format.space_after = Pt(9 if k == len(val) - 1 else 2)
                q.paragraph_format.line_spacing = LINE
                write_inline(q, t, size=BODY_PT - 0.5)
                for r in q.runs:
                    r.font.color.rgb = MUTED
                    r.italic = not t.startswith('(')

        elif kind == 'code':
            p = para()
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(val)
            set_font(r, MONO)
            r.font.size = Pt(BODY_PT - 1)

        elif kind == 'img':
            # 영어 빌드는 fig_en/ 을 먼저 본다. 화면 캡처나 형태소 시트처럼
            # 글자가 없는 도판은 거기 없으므로 fig/ 로 떨어진다. 언어와 무관한
            # 그림을 두 벌 두지 않으려는 것이다.
            path = (DOC / val).resolve()
            if LANG == 'EN':
                alt = DOC / 'fig_en' / path.name
                if alt.exists():
                    path = alt.resolve()
            p = para()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.add_run().add_picture(str(path), **fit(path))

        elif kind == 'table':
            head, *body = val
            t = cell.add_table(rows=len(val), cols=len(head))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            borders(t)
            for ci, txt in enumerate(head):
                c = t.rows[0].cells[ci]
                shade(c, 'F2F2F2')
                write_inline(clear(c), txt, size=BODY_PT - 1, bold=True)
            for ri, row in enumerate(body, start=1):
                for ci, txt in enumerate(row[:len(head)]):
                    write_inline(clear(t.rows[ri].cells[ci]), txt, size=BODY_PT - 1)
            # 표 다음 문단이 붙지 않게 한 줄 띄운다.
            cell.add_paragraph().paragraph_format.space_after = Pt(6)


def fit(path: Path) -> dict:
    """가로가 넘치면 폭으로, 세로가 넘치면 높이로 맞춘다 (휴대폰 화면이 후자다)."""
    from PIL import Image
    w, h = Image.open(path).size
    return {'width': Inches(MAX_W)} if h / w <= MAX_H / MAX_W else {'height': Inches(MAX_H)}


# ── 본문 자르기 ────────────────────────────────────────────────────

def sections(md: str) -> dict[str, str]:
    parts = re.split(r'^## □ (.+)$', md, flags=re.M)[1:]
    out = {}
    for title, body in zip(parts[0::2], parts[1::2]):
        key = next((k for k in SLOTS if title.startswith(k)), None)
        if key is None:
            raise SystemExit(f'템플릿에 없는 절: {title}')
        out[key] = body.strip().rstrip('-').strip()
    missing = set(SLOTS) - set(out)
    if missing:
        raise SystemExit(f'본문에 없는 절: {missing}')
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', default='proposal_ko.md',
                    help='본문 마크다운 (docs/proposal/ 기준)')
    ap.add_argument('--out', default=None)
    args = ap.parse_args()
    src = DOC / args.src
    global LANG
    LANG = 'EN' if src.stem.endswith('_en') else 'KO'
    if args.out is None:
        args.out = str(DOC / f'SnapCare_H1_Final_Proposal_{LANG}.docx')

    if not TEMPLATE.exists():
        raise SystemExit(f'템플릿을 찾지 못했다: {TEMPLATE}')

    d = docx.Document(str(TEMPLATE))
    body = sections(src.read_text())

    header = {**HEADER, **(HEADER_EN if LANG == 'EN' else {})}
    for (ri, ci), text in header.items():
        write_inline(clear(d.tables[0].rows[ri].cells[ci]), text, size=BODY_PT)

    for key, (ti, ri) in SLOTS.items():
        render(d.tables[ti].rows[ri].cells[0], body[key])
        print(f'  {key}')

    d.save(args.out)
    size = Path(args.out).stat().st_size / 1e6
    print(f'\n{args.out}  ({size:.1f} MB)')


if __name__ == '__main__':
    main()
