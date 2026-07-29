"""
도표의 원자료와 네이티브 엑셀 차트를 담은 통합문서를 만든다.

문서에 들어가는 그림은 charts.py 가 자동으로 만든다 — 수치가 바뀔 때마다 손으로
여덟 장을 다시 뽑는 것은 실수가 끼기 좋고, 실제로 이 문서는 인용을 대조하며 값이
열두 번 바뀌었다.

이 파일은 그 자동 경로를 대신하는 것이 아니라 **곁에 두는 것**이다. 발표 자료에
쓰거나 직접 손보고 싶을 때 열면 되고, 시트마다 자료와 차트가 함께 있어 어느
수치가 어느 그림에서 왔는지 바로 보인다.

    python3 docs/proposal/tools/build_xlsx.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference, Series
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill

OUT = Path(__file__).resolve().parent.parent / 'SnapCare_도표_원자료.xlsx'

HEAD = Font(bold=True, color='FFFFFF', size=10)
HEAD_BG = PatternFill('solid', fgColor='44546A')
TITLE = Font(bold=True, size=12)

# (시트 이름, 그림 번호, 제목, 머리행, 자료, 출처)
SHEETS = [
    ('그림1_문해율', '그림 1', '사하라이남 아프리카의 문해율과 비문해 인구',
     ['연도', '비문해 성인 (백만 명)', '문해율 (%)'],
     [[2015, 196, 65], [2024, 225, 69]],
     'UNESCO, International Literacy Day 2025 Factsheet'),

    ('그림2_건강문해력', '그림 2', '건강정보이해능력이 충분하지 않은 비율 (한국)',
     ['집단', '비율 (%)'],
     [['65세 이상', 59.7], ['소득 하위층', 45.6], ['전체 성인', 39.6], ['소득 상위층', 33.7]],
     '질병관리청, 국민건강영양조사 제9기 2차년도, 2023'),

    ('그림3_조제오류', '그림 3', '조제 오류의 유형별 분포 (약사가 발견한 155건)',
     ['유형', '건수', '비율 (%)'],
     [['누락', 51, 32.9], ['용량 오류', 44, 28.4], ['시각 오류', 36, 23.2], ['나머지', 24, 15.5]],
     'Anacleto 등, Clinics 62(3):243–250, 2007'),

    ('그림4_스마트폰', '그림 4', '연령대별 스마트폰 보유율 (미국, 2025)',
     ['연령대', '보유율 (%)'],
     [['50세 미만', 97], ['50~64세', 90], ['65세 이상', 78]],
     'Pew Research Center, Mobile Fact Sheet, 2025'),

    ('그림5_픽토그램', '그림 5', '의약 픽토그램이 이해도 기준을 통과한 비율',
     ['연구', '통과율 (%)', '적용 기준'],
     [['Gutierrez (2022) · USP+FIP 108종', 15.7, 'ANSI 85%'],
      ['Dowse & Ehlers (2001) · USP-DI 23종', 47.8, 'ANSI 85%'],
      ['같은 46명 · 현지 제작 23종', 87.0, 'ANSI 85%'],
      ['Dowse (2023) · 현지 제작 10종', 80.0, 'ISO 66.7%'],
      ['본 연구 SIGIL · 78종 (기계 판정)', 69.2, '의미 일치']],
     '기준이 달라 같은 자로 잰 값이 아니다'),

    ('그림15_형성평가', '그림 15', '형성 평가 — 수정 전후 (28회 시행)',
     ['항목', '수정 전', '수정 후'],
     [['여섯 항목 모두 만족', 0, 27], ['픽토그램 0장으로 끝남', 8, 0]],
     '표본 처방전 × 사용자 조건 28회'),

    ('그림17_모델오차', '그림 17', '영양 추정 모델의 항목별 오차 (n = 507)',
     ['항목', 'MAE', 'MAE%'],
     [['질량', 32.97, 16.61], ['열량', 58.12, 22.75], ['탄수화물', 6.02, 30.60],
      ['단백질', 5.41, 30.96], ['지방', 4.21, 32.82]],
     'Nutrition5k 공식 테스트 507접시 · 직접 채점'),
]


def build():
    wb = Workbook()
    wb.remove(wb.active)

    for name, figno, title, head, rows, src in SHEETS:
        ws = wb.create_sheet(name[:31])
        ws['A1'] = f'[{figno}] {title}'
        ws['A1'].font = TITLE
        ws['A2'] = f'출처 · {src}'
        ws['A2'].font = Font(size=9, color='808080')

        for c, h in enumerate(head, start=1):
            cell = ws.cell(row=4, column=c, value=h)
            cell.font, cell.fill = HEAD, HEAD_BG
            cell.alignment = Alignment(horizontal='center')
        for r, row in enumerate(rows, start=5):
            for c, v in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=v)

        ws.column_dimensions['A'].width = 34
        for col in 'BCD':
            ws.column_dimensions[col].width = 16

        last = 4 + len(rows)
        ch = BarChart()
        ch.type = 'col' if len(rows) <= 3 else 'bar'
        ch.style = 2
        ch.title = title
        ch.height, ch.width = 8.4, 16.5
        data = Reference(ws, min_col=2, max_col=min(3, len(head)),
                         min_row=4, max_row=last)
        cats = Reference(ws, min_col=1, min_row=5, max_row=last)
        ch.add_data(data, titles_from_data=True)
        ch.set_categories(cats)
        ch.dLbls = DataLabelList()
        ch.dLbls.showVal = True
        ws.add_chart(ch, f'A{last + 3}')

    # 첫 시트에 안내
    ws = wb.create_sheet('읽는 법', 0)
    for r, line in enumerate([
            'SnapCare 제안서 도표 원자료',
            '',
            '문서에 실리는 그림은 docs/proposal/tools/charts.py 가 자동으로 만든다.',
            '이 통합문서는 그 자동 경로를 대신하지 않는다 — 곁에 두는 것이다.',
            '',
            '· 발표 자료에 쓰거나 직접 손보고 싶을 때 연다',
            '· 시트마다 자료와 차트가 함께 있어 어느 수치가 어느 그림에서 왔는지 보인다',
            '· 수치를 여기서 고쳐도 문서의 그림은 바뀌지 않는다. 본문 값을 고쳤으면',
            '  charts.py 를 다시 돌려야 한다',
            '',
            '모든 수치는 docs/proposal/citations.md 에서 원문과 대조한 것이다.'], start=1):
        ws.cell(row=r, column=1, value=line)
    ws['A1'].font = Font(bold=True, size=13)
    ws.column_dimensions['A'].width = 88

    wb.save(OUT)
    print(f'{OUT.name}  (시트 {len(wb.sheetnames)}개)')


if __name__ == '__main__':
    build()
