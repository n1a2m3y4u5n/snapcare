"""설문 도판이 쓰는 수치를 엑셀에서 읽는다.

지금까지 chart_survey_who · chart_survey · chart_field_answers 가 수치를
코드에 박아 두고 있었다. 응답이 하나 늘었을 때 charts.py 를 돌려도 그림은
34명 그대로였고, 돌렸으니 됐다고 넘어갈 뻔했다. 자료에서 읽으면 그 일이
일어나지 않는다.
"""
import collections
import pathlib

XLSX = (pathlib.Path.home() /
        'Desktop/Claude/snapcare-survey/scratchpad/SnapCare_설문응답.xlsx')

# 응답자가 적은 나라를 대륙으로. 한 사람이 여러 나라를 적으면 첫 나라로 센다.
CONTINENT = {
    '대한민국': '아시아', '한국': '아시아', 'Indonesia': '아시아', 'India': '아시아',
    'Pakistan': '아시아', 'Malaysia': '아시아', 'Singapore': '아시아',
    'Bangladesh': '아시아', '일본': '아시아', '대만': '아시아',
    'Bulgaria': '유럽', 'Serbia': '유럽', 'Republic of Serbia': '유럽',
    'United Kingdom': '유럽', 'UK': '유럽', 'Portugal': '유럽', 'Ukraine': '유럽',
    'Nigeria': '아프리카', 'South Africa': '아프리카', 'Tunisia': '아프리카',
    'india': '아시아', 'Pakistan; Malaysia': '아시아',
    # 나라 대신 소속을 적은 응답이 하나 있다. 인도의 대학이다.
    'JSS College of Pharmacy': '아시아',
    'United States': '북미', 'USA': '북미', 'US': '북미',
    'Australia': '오세아니아',
    'Hungary': '유럽',
}

FIELD = {
    'Pharmacy / medication counselling': '약국·복약지도',
    'Nursing / public health': '간호·보건',
    'Other': '그 밖',
    'HCI / design': 'HCI·디자인',
}


def _rows():
    import openpyxl
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb['응답']
    hdr = [c.value for c in ws[1]]
    return [dict(zip(hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]


def survey_who():
    """(전체 인원, [(분야, 수)], [(대륙, 수)])."""
    rows = _rows()
    fc = collections.Counter(FIELD.get(r['b_field'], '그 외 셋') for r in rows)
    fields = [(k, fc[k]) for k in ('약국·복약지도', '간호·보건', '그 밖', 'HCI·디자인')
              if fc[k]]
    rest = sum(v for k, v in fc.items() if k == '그 외 셋')
    if rest:
        fields.append(('그 외 셋', rest))

    cc = collections.Counter()
    for r in rows:
        first = str(r['b_country'] or '').split(',')[0].strip()
        if first not in CONTINENT:
            # 조용히 아시아로 떨어뜨리면 유럽 한 명이 아시아가 된다. 실제로
            # "Republic of Serbia" 가 그랬다.
            raise SystemExit(f'대륙을 모르는 나라: {first!r}. CONTINENT 에 넣어라')
        cc[CONTINENT[first]] += 1
    order = ['아시아', '유럽', '아프리카', '북미', '오세아니아']
    conts = [(k, cc[k]) for k in order if cc[k]]
    return len(rows), fields, conts


def judgement():
    """판단 문항. (분모, [(답, 수)]). 판단 자격 다섯 구획만 센다."""
    QUALIFIED = ('Pharmacy / medication counselling', 'Nursing / public health',
                 'HCI / design', 'AAC / special education',
                 'Adult literacy education')
    rows = [r for r in _rows()
            if r['b_field'] in QUALIFIED and r.get('s1_alone')]
    c = collections.Counter(r['s1_alone'] for r in rows)
    order = ['Partly — some of it, not all', 'Yes, probably',
             'I have no basis to judge', 'No']
    return len(rows), [(k, c.get(k, 0)) for k in order]


def field_answers(field_key: str, question: str):
    """한 구획의 한 문항. [(답, 수)] 를 많은 순으로."""
    rows = [r for r in _rows() if r['b_field'] == field_key and r.get(question)]
    c = collections.Counter(r[question] for r in rows)
    return c.most_common()


if __name__ == '__main__':
    n, f, c = survey_who()
    print(f'응답 {n} · 분야 {f} · 대륙 {c}')
    d, j = judgement()
    print(f'판단 분모 {d} · {j}')
