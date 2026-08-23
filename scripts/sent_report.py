"""실제로 편지가 나간 사람 명단을 엑셀로 낸다.

두 자료를 토큰으로 잇는다.

    scratchpad/contacts/sent.jsonl               한 줄에 한 통. 보낸 기록 그 자체다.
    scratchpad/contacts/SnapCare_전문가_명단.xlsx  누구인지. 발송대상 시트.

"명단"과 "보낸 기록"을 구분한다. 명단에 있어도 발송이 실패했으면 보낸 것이 아니고,
sent.jsonl 에 줄이 있으면 서버가 250 을 돌려준 것이다. 이 파일이 기준이다.

응답 여부는 설문 덤프의 recipient 토큰으로 판정한다. 덤프 경로를 주지 않으면 그 열은
비워 둔다 — 없는 것을 "응답 안 함"으로 적으면 안 되기 때문이다.

    python3 scripts/sent_report.py [설문덤프디렉터리] [출력.xlsx]
"""
import json
import pathlib
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent
CONTACTS = ROOT / 'scratchpad' / 'contacts'
SENT = CONTACTS / 'sent.jsonl'
ROSTER = CONTACTS / 'SnapCare_전문가_명단.xlsx'

HEAD_FILL = PatternFill('solid', fgColor='40498C')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)
BODY = Font(size=10)
WRAP = Alignment(vertical='top', wrap_text=True)
TOP = Alignment(vertical='top')


def read_sent() -> list[dict]:
    rows = []
    for line in SENT.read_text(encoding='utf-8').splitlines():
        line = line.strip()
        if line:
            rows.append(json.loads(line))
    return rows


def read_roster() -> dict[str, dict]:
    """수신자 토큰 → 명단 한 행."""
    wb = load_workbook(ROSTER, read_only=True, data_only=True)
    ws = wb['발송대상']
    it = ws.iter_rows(values_only=True)
    header = [str(c or '').replace('\n', ' ').strip() for c in next(it)]
    idx = {name: i for i, name in enumerate(header)}
    out = {}
    for row in it:
        tok = str(row[idx['수신자 토큰']] or '').strip()
        if tok:
            out[tok] = {k: row[i] for k, i in idx.items()}
    wb.close()
    return out


def responded(spec) -> dict[str, str]:
    """수신자 토큰 → 접수 시각. 쉼표로 이은 여러 폴더를 받는다.

    웹 응답과 메일 응답이 서로 다른 폴더에 있어서, 한쪽만 읽으면 답한 사람이
    무응답으로 잡힌다.
    """
    if not spec:
        return {}
    out = {}
    for part in str(spec).split(','):
        d = pathlib.Path(part.strip())
        if not part.strip() or not d.is_dir():
            continue
        for f in sorted(d.glob('*.json')):
            j = json.loads(f.read_text(encoding='utf-8'))
            tok = str(j.get('recipient') or '').strip()
            if tok:
                out[tok] = j.get('receivedAt', '')
    return out


def style(ws, widths, freeze='A2'):
    for c in ws[1]:
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(vertical='center', wrap_text=True)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


def main() -> None:
    dump = sys.argv[1] if len(sys.argv) > 1 else None
    out = pathlib.Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / 'scratchpad' / '발송명단.xlsx'

    sent = read_sent()
    ros = read_roster()
    resp = responded(dump)
    print(f'발송 기록 {len(sent)}통, 명단 {len(ros)}명, 응답 {len(resp)}건')

    wb = Workbook()

    # ── 발송 ──
    ws = wb.active
    ws.title = '발송'
    cols = ['순번', '발송시각', '언어', '이름 / 기관', '소속·직위', '전문분야', '권역', '등급',
            '이메일', '수신자 토큰', '응답', '응답 접수시각', '담당 질문', '이 사람을 고른 이유']
    ws.append(cols)
    unmatched = 0
    for i, s in enumerate(sent, start=1):
        tok = s.get('token', '')
        r = ros.get(tok)
        if r is None:
            unmatched += 1
        r = r or {}
        ws.append([
            i, s.get('at', ''), s.get('lang', ''),
            r.get('이름 / 기관', ''), r.get('소속·직위', ''), r.get('전문분야', ''),
            r.get('권역', ''), r.get('등급', ''),
            s.get('email', ''), tok,
            'O' if tok in resp else ('' if not resp else ''),
            resp.get(tok, ''),
            r.get('담당 질문', ''), r.get('이 사람을 고른 이유', ''),
        ])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font, c.alignment = BODY, TOP
        row[3].alignment = row[4].alignment = row[12].alignment = row[13].alignment = WRAP
    style(ws, [6, 20, 6, 26, 30, 26, 8, 6, 32, 13, 6, 20, 34, 40])
    if unmatched:
        print(f'  경고: 명단에서 못 찾은 토큰 {unmatched}건')

    # ── 집계 ──
    ws = wb.create_sheet('집계')
    ws.append(['구분', '값', '발송', '응답', '응답률'])
    tok_by = {s['token']: s for s in sent}

    def block(title, keyfn):
        ws.append([title, '', '', '', ''])
        ws.cell(ws.max_row, 1).font = Font(bold=True, size=10)
        cnt, rep = Counter(), Counter()
        for s in sent:
            k = keyfn(s) or '(빈칸)'
            cnt[k] += 1
            if s['token'] in resp:
                rep[k] += 1
        for k, v in cnt.most_common():
            ws.append(['', k, v, rep[k], f'{rep[k] / v * 100:.1f}%' if resp else ''])
        ws.append(['', '합계', sum(cnt.values()), sum(rep.values()),
                   f'{sum(rep.values()) / len(sent) * 100:.1f}%' if resp else ''])
        ws.cell(ws.max_row, 2).font = Font(bold=True, size=10)
        ws.append([])

    block('언어', lambda s: s.get('lang'))
    block('권역', lambda s: (ros.get(s['token']) or {}).get('권역'))
    block('등급', lambda s: (ros.get(s['token']) or {}).get('등급'))
    block('발송 날짜(KST)', lambda s: str(s.get('at', ''))[:10])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if not c.font.bold:
                c.font = BODY
            c.alignment = TOP
    style(ws, [16, 30, 10, 10, 10], freeze='A2')
    ws.auto_filter.ref = None

    # ── 무응답 ──
    if resp:
        ws = wb.create_sheet('무응답')
        ws.append(['이름 / 기관', '소속·직위', '전문분야', '권역', '등급', '이메일', '언어'])
        for s in sent:
            if s['token'] in resp:
                continue
            r = ros.get(s['token']) or {}
            ws.append([r.get('이름 / 기관', ''), r.get('소속·직위', ''), r.get('전문분야', ''),
                       r.get('권역', ''), r.get('등급', ''), s.get('email', ''), s.get('lang', '')])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font, c.alignment = BODY, TOP
            row[0].alignment = row[1].alignment = WRAP
        style(ws, [26, 32, 28, 8, 6, 32, 6])

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print('저장:', out)
    for s in wb.worksheets:
        print(f'  {s.title}: {s.max_row - 1} 행')


if __name__ == '__main__':
    main()
